import io
import os
import re
from typing import Union
import pypdf
import docx

def format_pdf_text_spacing(text: str) -> str:
    """Format PDF text spacing to fix glued words, symbols, and years without corrupting emails or URLs."""
    lines = text.split("\n")
    cleaned_lines = []

    for line in lines:
        l_clean = line.strip()
        if not l_clean:
            continue

        parts = l_clean.split()
        new_parts = []
        for p in parts:
            if "@" in p or "http" in p or "linkedin" in p or "github" in p:
                new_parts.append(p)
            else:
                p_fmt = re.sub(r'([a-zA-Z])&([a-zA-Z])', r'\1 & \2', p)
                p_fmt = re.sub(r'([a-zA-Z]),([a-zA-Z])', r'\1, \2', p_fmt)
                p_fmt = re.sub(r'([a-zA-Z])((?:19|20)\d{2})', r'\1 \2', p_fmt)
                p_fmt = re.sub(r'([a-z])([A-Z])', r'\1 \2', p_fmt)
                p_fmt = re.sub(r'([A-Z]{2,})([A-Z][a-z])', r'\1 \2', p_fmt)
                new_parts.append(p_fmt)

        cleaned_lines.append(" ".join(new_parts))

    full = "\n".join(cleaned_lines)

    replacements = {
        "Instituteof": "Institute of",
        "Full-StackDeveloper": "Full-Stack Developer",
        "ScienceandEngineering": "Science and Engineering",
        "ComputerScience": "Computer Science",
        "TamilNadu": "Tamil Nadu",
        "MatricHr": "Matric Hr",
        "HrSec": "Hr Sec",
        "StateManagement": "State Management",
        "StylingLibrary": "Styling Library",
        "DeploymentPlatform": "Deployment Platform"
    }

    for k, v in replacements.items():
        full = full.replace(k, v)

    # Filter out template placeholders e.g. [Add 1-2 bullets...]
    full = re.sub(r'●\s*\[Add.*?\].*?\n', '', full)
    full = re.sub(r'\[.*?confirm.*?\]', '', full)
    full = re.sub(r'\[.*?library.*?\]', '', full)
    full = re.sub(r'\[.*?platform.*?\]', '', full)

    return full.strip()

def extract_pdf_text(file_source: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extract plain text from PDF using layout mode and single-character line reconstruction.
    """
    try:
        if isinstance(file_source, (bytes, bytearray)):
            reader = pypdf.PdfReader(io.BytesIO(file_source))
        elif isinstance(file_source, str) and os.path.exists(file_source):
            reader = pypdf.PdfReader(file_source)
        elif hasattr(file_source, 'read'):
            reader = pypdf.PdfReader(file_source)
        else:
            raise ValueError("Invalid PDF file source specified")

        page_texts = []
        for page in reader.pages:
            try:
                txt = page.extract_text(extraction_mode="layout")
            except Exception:
                txt = page.extract_text()

            if not txt:
                txt = page.extract_text() or ""

            # Check if PDF output has broken single-character lines
            lines = [l.strip() for l in txt.split("\n") if l.strip()]
            if len(lines) > 200 and sum(1 for l in lines if len(l) <= 2) / len(lines) > 0.4:
                reconstructed_words = []
                cur_word = []
                for l in lines:
                    if len(l) == 1:
                        cur_word.append(l)
                    else:
                        if cur_word:
                            reconstructed_words.append("".join(cur_word))
                            cur_word = []
                        reconstructed_words.append(l)
                if cur_word:
                    reconstructed_words.append("".join(cur_word))
                txt = " ".join(reconstructed_words)

            page_texts.append(txt)

        raw_full_text = "\n".join(page_texts).strip()
        formatted_text = format_pdf_text_spacing(raw_full_text)
        return formatted_text

    except Exception as e:
        raise ValueError(f"Failed to parse PDF document: {str(e)}")

def extract_docx_text(file_source: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extract plain text from a Word DOCX file path or binary byte buffer using python-docx.
    """
    full_text = []
    try:
        if isinstance(file_source, (bytes, bytearray)):
            doc = docx.Document(io.BytesIO(file_source))
        elif isinstance(file_source, str) and os.path.exists(file_source):
            doc = docx.Document(file_source)
        elif hasattr(file_source, 'read'):
            doc = docx.Document(file_source)
        else:
            raise ValueError("Invalid DOCX file source specified")

        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text.append(row_text)
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")

    return "\n".join(full_text).strip()

def extract_resume_text(file_source: Union[str, bytes, io.BytesIO], file_type: str) -> str:
    """
    Unified entry point for text extraction from PDF or DOCX files.
    """
    ft_lower = file_type.lower().strip()
    if "pdf" in ft_lower or ft_lower.endswith(".pdf"):
        return extract_pdf_text(file_source)
    elif "docx" in ft_lower or "openxmlformats" in ft_lower or ft_lower.endswith(".docx"):
        return extract_docx_text(file_source)
    else:
        raise ValueError(f"Unsupported file type '{file_type}'. Only PDF and DOCX files are allowed.")
